import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def is_cancelado(nome: str) -> bool:
    if not isinstance(nome, str):
        return False
    n = nome.strip().lower()
    return n == "cancelled" or n.endswith("(cancelled)")


def is_taxa_entrega(nome: str) -> bool:
    if not isinstance(nome, str):
        return False
    n = nome.strip().lower()
    return any(kw in n for kw in ("delivery fee", "shipping fee", "delivery", "shipping"))


def deve_excluir(numero_peca, nome: str) -> bool:
    if pd.isna(numero_peca) or str(numero_peca).strip() == "":
        return True
    if is_cancelado(nome):
        return True
    if is_taxa_entrega(nome):
        return True
    return False

def parse_system_report(caminho: str) -> pd.DataFrame:
    df = pd.read_excel(caminho, skiprows=4, dtype={"ID": str})
    df.columns = ["ID do Pedido", "Cliente", "Data", "Número de Peça", "Nome", "Preço"]
    df["ID do Pedido"] = df["ID do Pedido"].ffill()
    df["Cliente"] = df["Cliente"].ffill()
    mask = df.apply(
        lambda r: deve_excluir(r["Número de Peça"], r["Nome"]), axis=1
    )
    df = df[~mask].reset_index(drop=True)
    df["Data"] = pd.to_datetime(df["Data"], errors="coerce")
    return df[["ID do Pedido", "Data", "Cliente", "Número de Peça", "Nome", "Preço"]]

def parse_ql_log(caminho: str) -> pd.DataFrame:
    raw = pd.read_excel(
        caminho,
        header=None,
        dtype={1: str}, 
    )
    COL_STATUS       = 0 
    COL_ID_PEDIDO    = 1  
    COL_CLIENTE      = 2  
    COL_DATA         = 3  
    COL_NUMERO_PECA  = 7 
    COL_NOME         = 8  
    COL_PRECO        = 9  

    registros = []
    id_pedido_atual = None
    cliente_atual   = None
    data_atual      = None

    for _, linha in raw.iterrows():
        id_bruto = linha[COL_ID_PEDIDO]

        if (
            isinstance(id_bruto, str)
            and re.fullmatch(r"\d+", id_bruto.strip())
            and pd.isna(linha[COL_CLIENTE])
            and pd.isna(linha[COL_NUMERO_PECA])
        ):
            id_pedido_atual = id_bruto.strip()
            cliente_atual   = None
            data_atual      = None
            continue

        if id_pedido_atual is None:
            continue

        nome        = linha[COL_NOME]
        numero_peca = linha[COL_NUMERO_PECA]
        preco       = linha[COL_PRECO]
        cliente     = linha[COL_CLIENTE]
        data        = linha[COL_DATA]

        if pd.notna(cliente) and str(cliente).strip() not in ("DISCOUNT",):
            cliente_atual = str(cliente).strip()
        if pd.notna(data):
            data_atual = data
        if pd.isna(nome) or str(nome).strip() == "":
            continue

        if deve_excluir(numero_peca, nome):
            continue

        registros.append({
            "ID do Pedido":   id_pedido_atual,
            "Data":           data_atual,
            "Cliente":        cliente_atual,
            "Número de Peça": str(numero_peca).strip() if pd.notna(numero_peca) else None,
            "Nome":           str(nome).strip(),
            "Preço":          preco if pd.notna(preco) else None,
        })

    df = pd.DataFrame(registros, columns=["ID do Pedido", "Data", "Cliente", "Número de Peça", "Nome", "Preço"])
    df["Data"] = pd.to_datetime(df["Data"], errors="coerce")
    return df

def escrever_excel(df: pd.DataFrame, caminho: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório de Vendas"

    fonte_cabecalho  = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    fill_cabecalho   = PatternFill("solid", start_color="2F5496")
    alinhar_centro   = Alignment(horizontal="center", vertical="center")
    alinhar_esq      = Alignment(horizontal="left",   vertical="center")
    borda_fina       = Side(style="thin", color="CCCCCC")
    borda_celula     = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)
    fill_alternado   = PatternFill("solid", start_color="EEF2FA")
    fonte_padrao     = Font(name="Arial", size=10)

    cabecalhos = ["ID do Pedido", "Data", "Cliente", "Número de Peça", "Nome", "Preço"]
    for col_idx, cab in enumerate(cabecalhos, start=1):
        celula = ws.cell(row=1, column=col_idx, value=cab)
        celula.font      = fonte_cabecalho
        celula.fill      = fill_cabecalho
        celula.alignment = alinhar_centro
        celula.border    = borda_celula

    ws.row_dimensions[1].height = 20

    for linha_idx, registro in enumerate(df.itertuples(index=False), start=2):
        fill_linha = fill_alternado if linha_idx % 2 == 0 else None
        data_val   = registro.Data.date() if pd.notna(registro.Data) else None

        valores = [
            registro._0,      
            data_val,         
            registro.Cliente, 
            registro._3,      
            registro.Nome,   
            registro.Preço,  
        ]

        for col_idx, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha_idx, column=col_idx, value=valor)
            celula.font   = fonte_padrao
            celula.border = borda_celula
            if fill_linha:
                celula.fill = fill_linha

            if col_idx == 6 and valor is not None:   
                celula.number_format = '#,##0.00'
                celula.alignment = alinhar_centro
            elif col_idx == 2:                       
                celula.number_format = "YYYY-MM-DD"
                celula.alignment = alinhar_centro
            elif col_idx in (1, 4):                  
                celula.alignment = alinhar_centro
            else:
                celula.alignment = alinhar_esq

    larguras = [14, 14, 24, 16, 34, 10]
    for i, larg in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larg

    ws.freeze_panes = "A2"

    wb.save(caminho)
    print(f"Arquivo salvo: {caminho}  ({len(df)} linhas)")

def main():
    import os

    base               = os.path.dirname(os.path.abspath(__file__))
    caminho_system     = os.path.join(base, "System Report.xlsx")
    caminho_ql         = os.path.join(base, "QL Log.xlsx")
    caminho_saida      = os.path.join(base, "Relatório de Vendas.xlsx")

    print("Processando System Report …")
    df_system = parse_system_report(caminho_system)
    print(f"  → {len(df_system)} linhas válidas")

    print("Processando QL Log …")
    df_ql = parse_ql_log(caminho_ql)
    print(f"  → {len(df_ql)} linhas válidas")

    consolidado = pd.concat([df_system, df_ql], ignore_index=True)
    consolidado.sort_values(["ID do Pedido", "Data"], inplace=True)
    consolidado.reset_index(drop=True, inplace=True)

    print(f"Total de linhas consolidadas: {len(consolidado)}")
    escrever_excel(consolidado, caminho_saida)


if __name__ == "__main__":
    main()
